"""Genera src/data/obras.json a partir del reporte Excel (hoja Obras).
Criterios acordados: MATIAS LOPEZ se cuenta una sola vez (feb-25) y ORMAY 1 no se incluye.
Uso: python scripts/export_data.py "ruta/Reporte_Obras_Instaladas_2025-2026.xlsx"
"""
import json, re, sys
from openpyxl import load_workbook

SRC = sys.argv[1] if len(sys.argv) > 1 else "C:/Users/Solfl/Desktop/Dashboard indicadores/Reporte_Obras_Instaladas_2025-2026.xlsx"
OUT = "src/data/obras.json"
MES = ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]


def clean(s):
    return " ".join(str(s or "").split())


def estructura_grupo(e):
    if "Baratec" in e or "BT " in e:
        return "Baratec (suelo)"
    if e.startswith("Perfil"):
        return "Cochera solar"
    if re.search("Triangular|Telesc|Regulable", e):
        return "Triangular / Telescópica / Regulable"
    return e


def sistema(s):
    s = clean(s)
    return "Híbrido" if s.lower().startswith(("hib", "híb")) else s


wb = load_workbook(SRC, data_only=True)
ws = wb["Obras"]
obras = []
for r in range(5, 300):
    nombre = ws.cell(row=r, column=2).value
    if not nombre:
        continue
    nombre = clean(nombre)
    if nombre.upper().startswith("ORMAY"):
        continue
    g = lambda c: ws.cell(row=r, column=c).value
    anio = int(g(1))
    meses = []
    for m, y, v in re.findall(r"([a-záéíóú]{3})-(\d\d) \(([\d\.,]+) kWp\)", str(g(3))):
        meses.append({"ym": f"20{y}-{MES.index(m) + 1:02d}", "kwp": float(v.replace(".", "").replace(",", "."))})
    inst = float(g(6) or 0)
    if nombre.upper() == "MATIAS LOPEZ":  # duplicado en la hoja PI 2025: se cuenta una vez
        meses = [meses[0]]
        inst = meses[0]["kwp"]
    estr = clean(g(9)).replace("Estructura ", "")
    obras.append({
        "id": f"{nombre}|{anio}",
        "obra": nombre,
        "anio": anio,
        "meses": meses,
        "kwp": round(inst, 1),
        "potObra": round(float(g(4) or 0), 1),
        "rango": clean(g(5)),
        "implantacion": clean(g(8)).replace("Cochera Solar", "Cochera solar"),
        "estructura": estructura_grupo(estr),
        "estructuraDetalle": estr,
        "cliente": clean(g(10)),
        "sistema": sistema(g(11)),
        "marca": clean(g(12)).upper(),
        "acople": g(13),
        "wp": g(15),
        "paneles": g(17),
    })

json.dump(obras, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(len(obras), "obras;", round(sum(o["kwp"] for o in obras), 1), "kWp")
for a in (2025, 2026):
    print(a, sum(1 for o in obras if o["anio"] == a), round(sum(o["kwp"] for o in obras if o["anio"] == a), 1))
